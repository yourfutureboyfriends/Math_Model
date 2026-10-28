"""
LBO → Excel: the deal as an investment-committee workbook — Summary (returns by case, sources & uses, the value
bridge), Operating model, Debt schedule, Credit & covenants, Sensitivities, and Assumptions (every input and its
source). Figures are the model's computed values (api/pe/lbo.py), laid out the way deal teams read them.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "0F233E"
HEAD = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF")
SUB = PatternFill("solid", fgColor="E8EEF6")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color=NAVY)
THIN = Border(top=Side(style="thin", color="B7C4D6"))
NUM, PCT, X = '#,##0.0;(#,##0.0);"–"', "0.0%", '0.00"x"'


def _hdr(ws, row: int, labels: List[str], col: int = 1) -> None:
    for j, v in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=v)
        c.fill, c.font = HEAD, HEAD_FONT
        c.alignment = Alignment(horizontal="center" if j else "left")


def _line(ws, row: int, label: str, values: List[Any], fmt: str = NUM, bold: bool = False, total: bool = False) -> None:
    a = ws.cell(row=row, column=1, value=label)
    if bold:
        a.font = BOLD
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=2 + j, value=v)
        c.number_format = fmt
        if bold:
            c.font = BOLD
        if total:
            c.border = THIN


def _widths(ws, first: int = 34, rest: int = 13, n: int = 12) -> None:
    ws.column_dimensions["A"].width = first
    for j in range(2, n + 2):
        ws.column_dimensions[get_column_letter(j)].width = rest


def build(result: Dict[str, Any], sources: Optional[Dict[str, str]] = None, unit: float = 1e6) -> bytes:
    """The workbook bytes for an lbo.analyse() result; amounts shown in millions (unit) of the deal currency."""
    inp, base = result["inputs"], result["cases"]["base"]
    ccy, u = base.get("currency") or "", unit
    name = inp.get("name") or "Deal"
    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    _widths(ws, 38, 15, 6)
    ws["A1"], ws["A1"].font = f"{name} — LBO summary", TITLE
    ws["A2"] = f"{ccy} millions · base case exit in year {base['exit_year']} · modelled returns are analysis, not a track record"
    _hdr(ws, 4, ["Returns", "Base", "Upside", "Downside"])
    for i, (lab, key, fmt) in enumerate([("Sponsor IRR", "sponsor_irr", PCT), ("Money multiple (MOIC)", "moic", X)]):
        _line(ws, 5 + i, lab, [result["cases"][c]["returns"][key] for c in ("base", "upside", "downside")], fmt)
    _line(ws, 7, "Probability-weighted IRR", [result["weighted"]["irr"]], PCT, bold=True)
    me = result["max_entry"]
    _line(ws, 8, f"Max entry EV/EBITDA at a {me['hurdle']:.0%} IRR", [me.get("multiple")], X)
    _line(ws, 9, "EBITDA cushion to covenant breach", [result["cushion"]["to_covenant_breach"]], PCT)
    _hdr(ws, 11, ["Sources", f"{ccy}m", "x EBITDA", "% total"])
    r = 12
    tu = base["total_uses"]
    for k, v in base["sources"].items():
        _line(ws, r, k, [v / u, v / base["entry"]["ebitda"], v / tu], NUM)
        ws.cell(row=r, column=3).number_format, ws.cell(row=r, column=4).number_format = X, PCT
        r += 1
    _line(ws, r, "Total sources", [tu / u], NUM, bold=True, total=True)
    r += 2
    _hdr(ws, r, ["Uses", f"{ccy}m", "% total"])
    r += 1
    labels = {"equity_purchase": "Purchase of equity", "refinance_net_debt": "Refinance net debt", "transaction_fees": "Transaction fees",
              "financing_fees": "Financing fees", "cash_to_balance_sheet": "Cash to balance sheet"}
    for k, v in base["uses"].items():
        _line(ws, r, labels.get(k, k), [v / u, v / tu])
        ws.cell(row=r, column=3).number_format = PCT
        r += 1
    _line(ws, r, "Total uses", [tu / u], NUM, bold=True, total=True)
    r += 2
    _hdr(ws, r, ["Value creation (equity gain)", f"{ccy}m"])
    r += 1
    for k, lab in (("ebitda_growth", "EBITDA growth"), ("multiple", "Multiple expansion"), ("deleveraging", "Debt paydown / cash"),
                   ("fees_and_costs", "Fees and costs"), ("mip", "Management incentive plan"), ("total", "Total equity gain")):
        _line(ws, r, lab, [base["bridge"][k] / u], NUM, bold=k == "total", total=k == "total")
        r += 1

    # Operating model
    sch = base["schedule"]
    years = [f"Year {x['year']}" for x in sch]
    ws = wb.create_sheet("Operating model")
    _widths(ws, 34, 12, len(sch) + 1)
    ws["A1"], ws["A1"].font = f"{name} — operating model ({ccy}m, base case)", TITLE
    _hdr(ws, 3, ["", "LTM"] + years)
    rows = [("Revenue", "revenue", NUM, True), ("Growth", "growth", PCT, False), ("EBITDA", "ebitda", NUM, True),
            ("EBITDA margin", "margin", PCT, False), ("D&A", "dna", NUM, False), ("EBIT", "ebit", NUM, True),
            ("Cash interest", "cash_interest", NUM, False), ("PIK interest", "pik_interest", NUM, False),
            ("Financing-fee amortisation", "fee_amortisation", NUM, False), ("Tax", "tax", NUM, False),
            ("Net income", "net_income", NUM, True), ("Capex", "capex", NUM, False), ("Change in working capital", "change_nwc", NUM, False),
            ("Levered free cash flow", "fcf", NUM, True)]
    for i, (lab, k, fmt, b) in enumerate(rows):
        ltm = {"revenue": base["entry"]["revenue"], "ebitda": base["entry"]["ebitda"],
               "margin": base["entry"]["ebitda"] / base["entry"]["revenue"]}.get(k)
        vals = [ltm / u if (ltm is not None and fmt == NUM) else ltm] + [x[k] / u if fmt == NUM else x[k] for x in sch]
        _line(ws, 4 + i, lab, vals, fmt, bold=b)

    # Debt schedule
    ws = wb.create_sheet("Debt schedule")
    _widths(ws, 34, 12, len(sch) + 1)
    ws["A1"], ws["A1"].font = f"{name} — debt schedule ({ccy}m, closing balances)", TITLE
    _hdr(ws, 3, ["", "Close"] + years)
    r = 4
    for t in base["tranches"]:
        _line(ws, r, t["name"], [t["amount"] / u] + [x["balances"].get(t["name"], 0) / u for x in sch])
        _line(ws, r + 1, "   rate", [None] + [x["rates"].get(t["name"]) for x in sch], PCT)
        r += 2
    _line(ws, r, "Revolver", [0] + [x["revolver"] / u for x in sch])
    _line(ws, r + 1, "Total debt", [base["debt_raised"] / u] + [x["debt"] / u for x in sch], NUM, bold=True, total=True)
    _line(ws, r + 2, "Cash", [inp.get("min_cash", 0) / u] + [x["cash"] / u for x in sch])
    _line(ws, r + 3, "Net debt", [base["entry"]["net_debt_close"] / u] + [x["net_debt"] / u for x in sch], NUM, bold=True)
    _line(ws, r + 5, "Mandatory repayment", [None] + [x["mandatory"] / u for x in sch])
    _line(ws, r + 6, "Cash sweep", [None] + [x["sweep"] / u for x in sch])
    _line(ws, r + 7, "Revolver draw / (repay)", [None] + [(x["revolver_draw"] - x["revolver_repay"]) / u for x in sch])
    _line(ws, r + 8, "Base rate", [None] + [x["base_rate"] for x in sch], PCT)

    # Credit & covenants
    ws = wb.create_sheet("Credit & covenants")
    _widths(ws, 34, 12, len(sch) + 1)
    ws["A1"], ws["A1"].font = f"{name} — credit statistics and covenant tests", TITLE
    _hdr(ws, 3, [""] + years)
    for i, (lab, k, fmt) in enumerate([("Total debt / EBITDA", "leverage", X), ("Senior debt / EBITDA", "senior_leverage", X),
                                       ("Net debt / EBITDA", "net_leverage", X), ("EBITDA / cash interest", "interest_cover", X),
                                       ("(EBITDA − capex) / debt service", "fccr", X), ("FCF / debt", "fcf_to_debt", PCT)]):
        _line(ws, 4 + i, lab, [x[k] for x in sch], fmt)
    cov = base["covenants"]
    _line(ws, 11, "Covenant: max net leverage", [c["max_net_leverage"] for c in cov], X, bold=True)
    _line(ws, 12, "   headroom (EBITDA can fall)", [c["leverage_headroom"] for c in cov], PCT)
    _line(ws, 13, "Covenant: min interest cover", [c["min_interest_cover"] for c in cov], X, bold=True)
    for j, c in enumerate(cov):
        ws.cell(row=14, column=2 + j, value="pass" if c["leverage_ok"] and c["cover_ok"] else "BREACH")
    ws.cell(row=14, column=1, value="Covenant test")

    # Sensitivities
    ws = wb.create_sheet("Sensitivities")
    _widths(ws, 26, 11, 7)
    ws["A1"], ws["A1"].font = f"{name} — sponsor IRR sensitivities (base case)", TITLE
    r = 3
    for key in ("entry_exit", "leverage_growth", "year_exit"):
        g = result["sensitivity"][key]
        ws.cell(row=r, column=1, value=f"{g['row_label']} ↓  ·  {g['col_label']} →").font = BOLD
        fmtc = PCT if key == "leverage_growth" else X
        for j, c in enumerate(g["cols"]):
            cell = ws.cell(row=r + 1, column=2 + j, value=c)
            cell.number_format, cell.fill, cell.font = ("+0.0%;-0.0%" if fmtc == PCT else X), HEAD, HEAD_FONT
        for i, rv in enumerate(g["rows"]):
            lab = ws.cell(row=r + 2 + i, column=1, value=rv)
            lab.number_format, lab.fill = ("0" if key == "year_exit" else X), SUB
            for j, v in enumerate(g["irr"][i]):
                ws.cell(row=r + 2 + i, column=2 + j, value=v).number_format = PCT
        r += len(g["rows"]) + 4

    # Assumptions
    ws = wb.create_sheet("Assumptions")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110
    ws["A1"], ws["A1"].font = f"{name} — assumptions and sources", TITLE
    r = 3
    for k, v in (sources or {}).items():
        ws.cell(row=r, column=1, value=k.capitalize()).font = BOLD
        ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True)
        r += 1
    r += 1
    for k in ("entry_multiple", "exit_multiple", "exit_year", "transaction_fee_pct", "financing_fee_pct", "min_cash", "base_rate",
              "tax_rate", "interest_cap_pct", "dna_pct", "capex_pct", "nwc_pct", "mip_pct", "exit_fee_pct", "hurdle_irr"):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=str(inp.get(k)))
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
