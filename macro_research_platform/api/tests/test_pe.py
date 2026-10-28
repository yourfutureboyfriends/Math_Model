"""Private equity: the LBO engine against hand-built deals, the fund mathematics against closed forms, the deal
pipeline, and the API end to end (temporary databases)."""
import copy
import io

import pytest

from api.pe import deals, fundmath as fm, lbo


def simple(**kw):
    """A deal small enough to check by hand: 20 EBITDA bought at 5× (EV 100) with one 60 loan at a fixed 10%, no
    fees, no tax, no growth, no capex, no working capital; exit after 5 years at 5×."""
    inp = {**lbo.default_inputs(),
           "revenue": 100.0, "ebitda": 20.0, "cash": 0.0, "debt": 0.0, "entry_multiple": 5.0, "transaction_fee_pct": 0.0,
           "financing_fee_pct": 0.0, "min_cash": 0.0, "base_rate": 0.0,
           "tranches": [{"name": "Loan", "kind": "tlb", "multiple": 3.0, "floating": False, "coupon": 0.10, "amort_pct": 0.0,
                         "tenor": 10, "sweep": False}],
           "revolver": {"commitment": 0.0}, "interest_on": "opening", "years": 5, "revenue_growth": 0.0, "ebitda_margin": 0.20,
           "dna_pct": 0.0, "capex_pct": 0.0, "nwc_pct": 0.0, "tax_rate": 0.0, "interest_cap_pct": None, "exit_year": 5,
           "exit_multiple": 5.0, "exit_fee_pct": 0.0, "mip_pct": 0.0}
    inp.update(kw)
    return inp


def test_sources_uses_and_returns_by_hand():
    r = lbo.run(simple())
    assert r["sponsor_equity"] == pytest.approx(40.0) and r["total_uses"] == pytest.approx(100.0)
    # no sweep: 14 of cash a year piles up (20 EBITDA − 6 interest); exit equity 100 − (60 − 70) = 110
    assert r["exit"]["equity_net"] == pytest.approx(110.0)
    assert r["returns"]["moic"] == pytest.approx(110 / 40)
    assert r["returns"]["sponsor_irr"] == pytest.approx((110 / 40) ** (1 / 5) - 1, abs=1e-9)


def test_cash_sweep_by_hand():
    r = lbo.run(simple(tranches=[{**simple()["tranches"][0], "sweep": True}]))
    # opening-balance interest: 60 → 46 → 30.6 → 13.66 → repaid in year 4 with 4.974 left; year 5 adds 20
    bal = [x["balances"]["Loan"] for x in r["schedule"]]
    assert bal[:3] == pytest.approx([46.0, 30.6, 13.66])
    assert bal[3] == pytest.approx(0.0) and r["schedule"][4]["cash"] == pytest.approx(4.974 + 20.0)
    assert r["exit"]["equity_net"] == pytest.approx(124.974)


def test_bridge_sums_to_the_equity_gain():
    r = lbo.run(lbo.default_inputs())
    gain = r["exit"]["equity_net"] - r["sponsor_equity"] - r["rollover_equity"]
    assert r["bridge"]["total"] == pytest.approx(gain, rel=1e-9)


def test_average_balance_interest_is_solved():
    r = lbo.run(simple(interest_on="average", tranches=[{**simple()["tranches"][0], "sweep": True}]))
    y1 = r["schedule"][0]
    assert y1["cash_interest"] == pytest.approx(0.10 * (60 + y1["balances"]["Loan"]) / 2, rel=1e-6)


def test_interest_cap_and_carry_forward():
    # 30% of 20 EBITDA = 6 deductible; a 9 interest bill leaves 3 a year carried forward
    r = lbo.run(simple(tax_rate=0.25, interest_cap_pct=0.30, tranches=[{**simple()["tranches"][0], "coupon": 0.15}]))
    y1 = r["schedule"][0]
    assert y1["interest_deductible"] == pytest.approx(6.0) and y1["interest_disallowed_cf"] == pytest.approx(3.0)
    assert y1["tax"] == pytest.approx((20 - 6) * 0.25)


def test_pik_mezzanine_accrues_and_revolver_funds_a_shortfall():
    mezz = {"name": "Mezz", "kind": "mezz", "multiple": 1.0, "floating": False, "coupon": 0.12, "pik": True, "amort_pct": 0, "tenor": 10}
    r = lbo.run(simple(tranches=simple()["tranches"] + [mezz]))
    assert r["schedule"][0]["balances"]["Mezz"] == pytest.approx(20 * 1.12)
    assert r["schedule"][4]["balances"]["Mezz"] == pytest.approx(20 * 1.12 ** 5)
    # capex above EBITDA in year 1 → the revolver draws the gap
    r2 = lbo.run(simple(capex_pct=[0.30, 0.0, 0.0, 0.0, 0.0], revolver={"commitment": 50.0, "spread": 0.0, "commitment_fee": 0.0}))
    assert r2["schedule"][0]["revolver_draw"] == pytest.approx(30 + 6 - 20)


def test_max_entry_hits_the_hurdle_exactly():
    inp = lbo.default_inputs()
    m = lbo.max_entry(inp, 0.20)
    q = {**inp, "entry_multiple": m["multiple"], "exit_multiple": m["exit_multiple"]}
    assert lbo.run(q)["returns"]["sponsor_irr"] == pytest.approx(0.20, abs=1e-6)


def test_analysis_runs_and_flags_overleverage():
    a = lbo.analyse(lbo.default_inputs())
    assert set(a["cases"]) == {"base", "upside", "downside"}
    assert a["cases"]["upside"]["returns"]["sponsor_irr"] > a["cases"]["base"]["returns"]["sponsor_irr"] > a["cases"]["downside"]["returns"]["sponsor_irr"]
    assert len(a["sensitivity"]["entry_exit"]["irr"]) == 5
    heavy = copy.deepcopy(lbo.default_inputs())
    heavy["tranches"][0]["multiple"] = 7.0
    assert any("leverage" in w for w in lbo.run(heavy)["warnings"])
    with pytest.raises(lbo.LboError):
        lbo.run({**lbo.default_inputs(), "ebitda": -5})


def test_xirr_and_multiples_match_closed_forms():
    assert fm.xirr([{"date": "2020-01-01", "amount": -100}, {"date": "2021-01-01", "amount": 110}]) == pytest.approx(1.1 ** (365 / 366) - 1, abs=1e-9)   # 2020 is a leap year
    p = fm.performance([{"date": "2020-01-01", "amount": -100}, {"date": "2023-01-01", "amount": 150}], nav=50, nav_date="2025-01-01")
    assert (p["dpi"], p["rvpi"], p["tvpi"]) == pytest.approx((1.5, 0.5, 2.0))


def test_ks_pme_by_hand():
    idx = {f"{y}-01-01": 100 * 1.1 ** (y - 2020) for y in range(2020, 2026)}
    p = fm.performance([{"date": "2020-01-01", "amount": -100}, {"date": "2023-01-01", "amount": 150}], nav=50, nav_date="2025-01-01", index=idx)
    assert p["pme"]["ks_pme"] == pytest.approx((150 * 1.1 ** 2 + 50) / (100 * 1.1 ** 5))
    assert "ln_pme_note" in p["pme"]                       # LN-PME undefined when the index portfolio goes negative


def test_waterfall_by_hand():
    w = fm.waterfall([{"date": "2020-01-01", "amount": -100}], [{"date": "2025-01-01", "amount": 200}])
    assert w["totals"]["preferred"] == pytest.approx(100 * (1.08 ** (1827 / 365) - 1))
    assert w["gp_carry"] == pytest.approx(20.0) and w["lp_total"] == pytest.approx(180.0)
    low = fm.waterfall([{"date": "2020-01-01", "amount": -100}], [{"date": "2025-01-01", "amount": 130}])
    assert low["gp_carry"] == 0 and low["totals"]["preferred"] == pytest.approx(30.0)
    half = fm.waterfall([{"date": "2020-01-01", "amount": -100}], [{"date": "2025-01-01", "amount": 200}], catch_up=0.5)
    assert half["gp_carry"] == pytest.approx(20.0)        # enough profit: a partial catch-up still reaches 20%


def test_pipeline_keeps_deals_and_logs_stages(tmp_path, monkeypatch):
    from api import funds
    monkeypatch.setattr(funds, "db_path", lambda fund_id=None: str(tmp_path / "fund.db"))
    d = deals.save("pm", lbo.default_inputs(), deals.summarise(lbo.analyse(lbo.default_inputs())), name="Project Alpha")
    assert d["stage"] == "Sourcing" and d["summary"]["irr"] is not None
    deals.set_stage("pm", d["id"], "Due diligence")
    deals.add_note("analyst", d["id"], "Management meeting went well.")
    got = deals.get(d["id"])
    assert got["stage"] == "Due diligence" and [x["kind"] for x in got["log"]][:2] == ["note", "stage"]
    with pytest.raises(ValueError):
        deals.set_stage("pm", d["id"], "Bogus")
    deals.delete(d["id"])
    assert deals.list_deals() == []


def test_excel_workbook_has_every_sheet():
    from openpyxl import load_workbook
    from api.pe import excel
    data = excel.build(lbo.analyse(lbo.default_inputs()), {"company": "test"}, unit=1.0)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Operating model", "Debt schedule", "Credit & covenants", "Sensitivities", "Assumptions"]


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    from fastapi.testclient import TestClient
    from api import funds, portfolio_store
    path = str(tmp_path_factory.mktemp("pe") / "pe.db")
    mp = pytest.MonkeyPatch()
    mp.setattr(portfolio_store, "_db_path", lambda: path)
    mp.setattr(funds, "db_path", lambda fund_id=None: path)
    import api.core.auth as auth
    mp.setattr(auth, "get_db_path", lambda: path)
    from api.core import accounts
    accounts.invalidate_cache()
    accounts.migrate()
    accounts.set_password("pm", "Test-Pm-Pass-2026!")
    from api.main import app
    with TestClient(app) as c:
        tok = c.post("/api/auth/login", json={"username": "pm", "password": "Test-Pm-Pass-2026!"}).json()["access_token"]
        c.headers.update({"Authorization": f"Bearer {tok}"})
        yield c
    mp.undo()
    accounts.invalidate_cache()


def test_api_runs_saves_and_exports(client):
    assert client.get("/api/v1/pe/lbo").status_code == 405
    inp = client.get("/api/v1/pe/defaults").json()["inputs"]
    r = client.post("/api/v1/pe/lbo", json={"inputs": inp})
    assert r.status_code == 200 and r.json()["cases"]["base"]["returns"]["sponsor_irr"] is not None
    assert client.post("/api/v1/pe/lbo", json={"inputs": {**inp, "ebitda": 0}}).status_code == 400
    d = client.post("/api/v1/pe/deals", json={"inputs": inp, "name": "Project Beta"}).json()
    assert d["name"] == "Project Beta" and d["summary"]["moic"] > 1
    assert client.post(f"/api/v1/pe/deals/{d['id']}/stage", json={"stage": "Screening"}).json()["stage"] == "Screening"
    assert [x["name"] for x in client.get("/api/v1/pe/deals").json()["deals"]] == ["Project Beta"]
    x = client.post("/api/v1/pe/lbo/excel", json={"inputs": inp})
    assert x.status_code == 200 and x.content[:2] == b"PK"
    w = client.post("/api/v1/pe/fund/waterfall", json={"contributions": [{"date": "2020-01-01", "amount": -100}],
                                                      "distributions": [{"date": "2025-01-01", "amount": 200}]}).json()
    assert w["gp_carry"] == pytest.approx(20.0)
