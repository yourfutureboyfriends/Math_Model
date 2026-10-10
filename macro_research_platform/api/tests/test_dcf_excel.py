"""The Excel export must compute the same value as the terminal's DCF. openpyxl doesn't
calculate formulas, so a tiny evaluator recalculates the workbook here."""
import io
import re

import pytest
from openpyxl import load_workbook

from api.tests.test_markets_functions import _inp


def _evaluate(wb, ref):
    cache = {}

    def cell(sheet, addr):
        key = (sheet, addr)
        if key in cache:
            return cache[key]
        v = wb[sheet][addr].value
        if isinstance(v, str) and v.startswith("="):
            v = expr(sheet, v[1:])
        cache[key] = v
        return v

    def expr(sheet, f):
        def rng(m):                                             # SUM(C15:L15) → row range
            c1, r1, c2, r2 = m.group(1), m.group(2), m.group(3), m.group(4)
            from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
            vals = [f"__c('{sheet}','{gl(k)}{r1}')" for k in range(ci(c1), ci(c2) + 1)]
            return "[" + ",".join(vals) + "]"
        f = re.sub(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", lambda m: "sum(" + rng(m) + ")", f)
        f = re.sub(r"IF\(([^,]+),([^,]+),([^)]+)\)", r"((\2) if (\1) else (\3))", f)
        f = f.replace("MAX(", "max(").replace("^", "**")
        f = re.sub(r"(Inputs|Model)!\$?([A-Z]+)\$?(\d+)", r"__c('\1','\2\3')", f)
        f = re.sub(r"(?<![A-Za-z_'])([A-Z]{1,2})(\d+)(?![\d'])", lambda m: f"__c('{sheet}','{m.group(1)}{m.group(2)}')", f)
        return eval(f, {"__c": cell, "sum": sum, "max": max})

    sheet, addr = ref.split("!")
    return cell(sheet, addr)


@pytest.mark.parametrize("kw", [dict(years=10), dict(years=5, mid_year=False), dict(years=25, growth=0.15),
                                dict(years=10, target_margin=0.30, include_leases=True)])
def test_excel_matches_model(kw):
    from api.marketdata import dcf, dcf_excel
    inp = _inp(leases=150.0, growth_y1=0.12, investments=80.0)
    v = dcf.value(inp, erp=0.05, **kw)
    data, keys = dcf_excel.build(inp, v)
    wb = load_workbook(io.BytesIO(data))
    assert _evaluate(wb, keys["ev"]) == pytest.approx(v["enterprise_value"], rel=1e-9)
    assert _evaluate(wb, keys["per_share"]) == pytest.approx(v["per_share"], rel=1e-9)
